import os
from pathlib import Path
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.units import inch
from reportlab.lib.colors import black, white, Color
from reportlab.pdfbase.pdfmetrics import stringWidth
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

ROOT = Path(os.environ.get('REPO_ROOT', '.')).resolve()
(ROOT/'downloads/templates').mkdir(parents=True, exist_ok=True)
(ROOT/'downloads/presentations').mkdir(parents=True, exist_ok=True)
(ROOT/'downloads/teacher-guides').mkdir(parents=True, exist_ok=True)
STUDENT = ROOT/'downloads/templates/IED_1_8_Engineering_Documentation_Challenge_Student_Packet.pdf'
TEACHER = ROOT/'downloads/teacher-guides/IED_1_8_Engineering_Documentation_Challenge_Teacher_Guide.pdf'
PPTX = ROOT/'downloads/presentations/IED_1_8_Engineering_Documentation_Challenge_Presentation.pptx'
LOGO = ROOT/'assets/img/lockwoodstem-concept-a1-black-transparent.png'
W,H = LETTER
M = 0.45*inch
DARK = Color(0.08,0.12,0.18)
LIGHT = Color(0.94,0.94,0.94)
GRAY = Color(0.85,0.85,0.85)

def wrap(text, font, size, width):
    words=text.split(); lines=[]; line=''
    for word in words:
        test=word if not line else line+' '+word
        if stringWidth(test,font,size)<=width: line=test
        else:
            if line: lines.append(line)
            line=word
    if line: lines.append(line)
    return lines

def wtext(c,text,x,y,width,size=8,leading=9.5,font='Helvetica',max_lines=None):
    lines=wrap(text,font,size,width)
    if max_lines: lines=lines[:max_lines]
    c.setFillColor(black); c.setFont(font,size)
    for line in lines:
        c.drawString(x,y,line); y-=leading
    return y

def header(c,title,subtitle,page):
    c.setFillColor(DARK); c.rect(0,H-.73*inch,W,.73*inch,fill=1,stroke=0)
    c.setFillColor(white); c.setFont('Helvetica-Bold',15); c.drawString(M,H-.29*inch,title)
    c.setFont('Helvetica',8.5); c.drawString(M,H-.50*inch,subtitle)
    c.setFont('Helvetica-Bold',8); c.drawRightString(W-M,H-.64*inch,f'PAGE {page}')
    if LOGO.exists(): c.drawImage(str(LOGO),W-1.85*inch,H-.60*inch,width=1.25*inch,height=.35*inch,preserveAspectRatio=True,mask='auto')

def info(c):
    y=H-.96*inch; c.setFillColor(black); c.setFont('Helvetica',8); c.setStrokeColor(black); c.setLineWidth(.6)
    c.drawString(M,y,'NAME:'); c.line(M+.42*inch,y-2,M+2.75*inch,y-2)
    c.drawString(M+3*inch,y,'PERIOD:'); c.line(M+3.58*inch,y-2,M+4.4*inch,y-2)
    c.drawString(M+4.68*inch,y,'DATE:'); c.line(M+5.05*inch,y-2,W-M,y-2)

def bar(c,y,title,right=''):
    h=.30*inch; c.setFillColor(LIGHT); c.setStrokeColor(black); c.setLineWidth(.6); c.rect(M,y-h,W-2*M,h,fill=1,stroke=1)
    c.setFillColor(black); c.setFont('Helvetica-Bold',10); c.drawString(M+.1*inch,y-.205*inch,title)
    if right: c.setFont('Helvetica',7.2); c.drawRightString(W-M-.1*inch,y-.205*inch,right)
    return y-h-.08*inch

def box(c,x,y,w,h,title):
    c.setFillColor(white); c.setStrokeColor(black); c.setLineWidth(.65); c.roundRect(x,y-h,w,h,7,fill=1,stroke=1)
    c.setFont('Helvetica-Bold',8.5); c.setFillColor(black); c.drawString(x+.08*inch,y-.19*inch,title)

def line(c,x,y,label,width):
    c.setFont('Helvetica-Bold',7.4); c.drawString(x,y,label)
    lw=stringWidth(label,'Helvetica-Bold',7.4); c.setLineWidth(.4); c.line(x+lw+4,y-2,x+width,y-2)

def table(c,x,y,cols,row_h,heads,rows,fs=6):
    total=sum(cols); c.setStrokeColor(black); c.setLineWidth(.5); c.setFillColor(LIGHT); c.rect(x,y-row_h,total,row_h,fill=1,stroke=1)
    xx=x; c.setFillColor(black); c.setFont('Helvetica-Bold',fs)
    for i,(h,w) in enumerate(zip(heads,cols)):
        c.drawCentredString(xx+w/2,y-row_h*.67,h); xx+=w
        if i<len(cols)-1: c.line(xx,y-row_h,xx,y)
    yy=y-row_h
    for r in rows:
        c.setFillColor(white); c.rect(x,yy-row_h,total,row_h,fill=1,stroke=1)
        xx=x; c.setFillColor(black); c.setFont('Helvetica',fs)
        for i,w in enumerate(cols):
            if i<len(cols)-1: c.line(xx+w,yy-row_h,xx+w,yy)
            if i < len(r) and r[i]: c.drawString(xx+3,yy-row_h*.67,r[i])
            xx+=w
        yy-=row_h
    return yy

def sketch(c,x,y,w,h,title,note):
    box(c,x,y,w,h,title); c.setStrokeColor(Color(.88,.88,.88)); c.setLineWidth(.25)
    gx=x+.1*inch
    while gx<x+w-.08*inch: c.line(gx,y-.35*inch,gx,y-h+.25*inch); gx+=.25*inch
    gy=y-.35*inch
    while gy>y-h+.25*inch: c.line(x+.08*inch,gy,x+w-.08*inch,gy); gy-=.25*inch
    c.setFillColor(Color(.45,.45,.45)); c.setFont('Helvetica-Oblique',7); c.drawString(x+.1*inch,y-h+.12*inch,note)

def make_student():
    c=canvas.Canvas(str(STUDENT),pagesize=LETTER)
    header(c,'IED 1.8 | Engineering Documentation Challenge','Unit 1 Final Challenge | Student Packet',1); info(c)
    y=bar(c,H-1.20*inch,'MISSION BRIEF','Use evidence, not guessing')
    wtext(c,'You will document a small purpose-made assembly so another engineering team could understand, reproduce, inspect, and improve it. The physical assembly is your source of evidence. Your final package should communicate the design clearly without requiring a verbal explanation.',M+.1*inch,y,W-2*M-.2*inch,8.3,10)
    y-=.55*inch; box(c,M,y,3.4*inch,2.05*inch,'DRIVING QUESTION'); box(c,M+3.65*inch,y,W-2*M-3.65*inch,2.05*inch,'FINAL DELIVERABLE')
    wtext(c,'How can an engineer fully communicate an existing assembly so another person could reproduce it without guessing?',M+.12*inch,y-.42*inch,3.16*inch,10,13,'Helvetica-Bold')
    yy=y-.42*inch
    for item in ['isometric overview sketch','aligned front/top/right-side views','complete dimensions','measurement and tolerance evidence','joining/interface analysis','justified improvement idea','peer review and final revision check']:
        c.setFont('Helvetica',7.7); c.drawString(M+3.80*inch,yy,'- '+item); yy-=.20*inch
    y-=2.25*inch; y=bar(c,y,'CHALLENGE RULES','Keep your work neat enough for another team to use')
    for i,r in enumerate(['Measure the actual assembly before finalizing critical dimensions.','Do not dimension from memory or guess at hidden features.','Use proper line conventions, centerlines, hidden lines, and dimension placement.','Avoid repeated dimensions. Every important feature should be sized and located.','Use tolerances only where they help control function, fit, alignment, or assembly.','Your improvement must be based on evidence from the existing design.'],1):
        c.setFont('Helvetica-Bold',7.6); c.drawString(M+.1*inch,y,f'{i}.'); wtext(c,r,M+.32*inch,y,W-2*M-.42*inch,7.6,9,max_lines=2); y-=.28*inch
    y=bar(c,y,'ASSEMBLY SNAPSHOT','Complete this before detailed documentation')
    for lab in ['Assembly name or station ID:','Primary function of the assembly:','Main parts you can identify:','Most important interface or connection:']:
        line(c,M+.1*inch,y,lab,W-2*M-.2*inch); y-=.29*inch
    c.showPage()
    header(c,'IED 1.8 | Product Investigation','Observe the assembly before drawing it',2); info(c)
    y=bar(c,H-1.20*inch,'PARTS AND INTERFACES','Identify how the assembly is built')
    y=table(c,M+.05*inch,y,[1*inch,1.55*inch,1.55*inch,1.95*inch,1*inch],.36*inch,['Part ID','Part name/description','Material/process','Function in assembly','Qty'],[['A','','','',''],['B','','','',''],['C','','','',''],['D','','','','']],6.1)-.18*inch
    y=bar(c,y,'INTERFACE MAP','Where do parts touch, connect, align, or move?')
    y=table(c,M+.05*inch,y,[1*inch,1*inch,1.55*inch,1.45*inch,2*inch],.34*inch,['Interface','Parts','Connection','Function','Evidence'],[['1','','','',''],['2','','','',''],['3','','','',''],['4','','','','']],5.9)-.18*inch
    y=bar(c,y,'ISOMETRIC OVERVIEW SKETCH','Show the full assembly in one clear view'); sketch(c,M+.05*inch,y,W-2*M-.1*inch,2.92*inch,'ISOMETRIC SKETCH SPACE','Label major parts and important interfaces.'); c.showPage()
    header(c,'IED 1.8 | Multiview Documentation','Create aligned views using proper line conventions',3); info(c)
    y=bar(c,H-1.20*inch,'VIEW PLANNING','Choose views that best communicate the assembly')
    for lab in ['Front view chosen because:','Top view should show:','Right-side view should show:','Hidden or center lines needed for:']:
        line(c,M+.1*inch,y,lab,W-2*M-.2*inch); y-=.27*inch
    y=bar(c,y-.03*inch,'MULTIVIEW SKETCH SPACE','Align Front, Top, and Right-side views')
    ww=(W-2*M-.2*inch)/2; sketch(c,M+.05*inch,y,ww,2.18*inch,'FRONT VIEW','Use object, hidden, and center lines.'); sketch(c,M+.15*inch+ww,y,ww,2.18*inch,'RIGHT-SIDE VIEW','Keep views aligned where possible.'); y-=2.33*inch; sketch(c,M+.05*inch,y,W-2*M-.1*inch,2.18*inch,'TOP VIEW','Line up features with the front view.'); c.showPage()
    header(c,'IED 1.8 | Measurement and Dimensioning','Size and locate the features another team needs',4); info(c)
    y=bar(c,H-1.20*inch,'MEASUREMENT RECORD','Measure before final dimension decisions')
    y=table(c,M+.05*inch,y,[1.55*inch,1.05*inch,1.05*inch,1.05*inch,1.05*inch,1.3*inch],.31*inch,['Feature measured','Tool used','Trial 1','Trial 2','Trial 3','Reported value'],[['','','','','',''] for _ in range(8)],5.8)-.16*inch
    y=bar(c,y,'CRITICAL DIMENSIONS AND TOLERANCES','Choose dimensions that control function or fit')
    y=table(c,M+.05*inch,y,[1.35*inch,1.15*inch,1.15*inch,1.45*inch,2*inch],.38*inch,['Critical feature','Nominal value','Tolerance','Why it matters','What happens if wrong?'],[['','','','',''] for _ in range(5)],5.5)-.12*inch
    c.setFont('Helvetica-Bold',7.5); c.drawString(M+.05*inch,y,'Dimensioning check:'); c.setFont('Helvetica',7.2); c.drawString(M+1.35*inch,y,'Every needed feature is sized and located. No dimensions are repeated unnecessarily.'); c.showPage()
    header(c,'IED 1.8 | Interface and Joining Analysis','Explain how the assembly works, not just what it looks like',5); info(c)
    y=bar(c,H-1.20*inch,'JOINING / INTERFACE ANALYSIS','Use the decision factors from Lesson 1.7')
    y=table(c,M+.05*inch,y,[1.25*inch,1.35*inch,1.1*inch,1.05*inch,1.6*inch,1*inch],.37*inch,['Connection','Method','Removable?','Function','Best reason','Concern'],[['','','','','',''] for _ in range(5)],5.4)-.18*inch
    y=bar(c,y,'ENGINEERING IMPROVEMENT PROPOSAL','Base your idea on evidence from the existing assembly')
    for lab in ['Problem or opportunity observed:','Evidence from measurements, sketches, or interface analysis:','Proposed design change:','How the change improves function, assembly, serviceability, or manufacturability:']:
        line(c,M+.1*inch,y,lab,W-2*M-.2*inch); y-=.31*inch
    c.setFont('Helvetica-Bold',7.4); c.drawString(M+.1*inch,y,'Sketch of proposed improvement:'); y-=.1*inch; sketch(c,M+.05*inch,y,W-2*M-.1*inch,1.45*inch,'IMPROVEMENT SKETCH','Label what changed.'); c.showPage()
    header(c,'IED 1.8 | Peer Review and Final Check','Could another team reproduce this without guessing?',6); info(c)
    y=bar(c,H-1.20*inch,'PEER ENGINEERING REVIEW','Reviewer completes this section')
    for i,chk in enumerate(['The assembly purpose is clear.','Major parts are identified and labeled.','Views are aligned and use correct line types.','Features are sized and located without obvious missing dimensions.','Critical dimensions and tolerances are reasonable.','Joining/interface analysis explains how the assembly works.','The improvement proposal is supported by evidence.'],1):
        c.setFont('Helvetica',7.6); c.drawString(M+.12*inch,y,f'{i}. {chk}'); c.rect(W-M-.7*inch,y-3,.1*inch,.1*inch); y-=.23*inch
    y-=.05*inch
    for lab in ['Reviewer name:','One strength of this documentation package:','One specific revision needed before final submission:']:
        line(c,M+.1*inch,y,lab,W-2*M-.2*inch); y-=.30*inch
    y=bar(c,y-.05*inch,'FINAL SUBMISSION RUBRIC','Use this before turning in your package')
    rows=[['Sketching and views','20','Clear isometric sketch; aligned orthographic views; correct line conventions'],['Dimensioning','20','Complete sizes and locations; proper symbols; no repeated or missing key dimensions'],['Measurement and tolerances','20','Accurate records; critical dimensions; tolerances tied to function or fit'],['Joining/interface analysis','15','Connection method, parts involved, function, and tradeoffs explained'],['Improvement proposal','15','Evidence-based change with clear sketch and justification'],['Professional quality','10','Neat, organized, labeled, readable, and complete']]
    y=table(c,M+.05*inch,y,[2*inch,1.1*inch,3.9*inch],.36*inch,['Category','Points','What strong work includes'],rows,5.9)-.2*inch
    line(c,M+.1*inch,y,'Final revision completed by:',W-2*M-.2*inch); c.save()

def make_teacher():
    c=canvas.Canvas(str(TEACHER),pagesize=LETTER)
    header(c,'IED 1.8 | Teacher Guide','Engineering Documentation Challenge | Setup and Facilitation',1)
    y=bar(c,H-1.02*inch,'LESSON OVERVIEW','Recommended length: 2-3 class periods')
    y=wtext(c,'Students reverse-engineer a small purpose-made assembly and produce a technical documentation package that integrates isometric sketching, multiview drawing, line conventions, dimensioning, measurement, tolerances, joining methods, and evidence-based improvement.',M+.1*inch,y,W-2*M-.2*inch,8.2,10)-.15*inch
    y=bar(c,y,'SUGGESTED ASSEMBLY: AEROSPACE SENSOR MOUNT MOCK-UP','Reusable and purpose-made')
    for name,desc in [('Base plate','3.000 x 1.500 x 0.250 in; two mounting holes; one attachment interface'),('Upright bracket','1.500 x 1.250 x 0.250 in vertical plate; sensor hole or slot; attaches to base'),('Spacer or locator block','0.500 x 0.500 x 0.250 in; creates offset or alignment feature'),('Joining hardware','Use screws, bolts/nuts, or tab/slot plus fastener depending on the interface you want students to analyze')]:
        c.setFont('Helvetica-Bold',7.8); c.drawString(M+.1*inch,y,name+':'); wtext(c,desc,M+1.55*inch,y,W-2*M-1.65*inch,7.6,9,max_lines=2); y-=.31*inch
    y=bar(c,y-.02*inch,'CLASSROOM FLOW','Keep most work student-driven')
    for day,desc in [('Day 1','Introduce challenge, identify parts/interfaces, complete assembly snapshot, isometric sketch, and view planning.'),('Day 2','Measure parts, create multiview documentation, add dimensions, identify critical dimensions, and assign tolerances.'),('Day 3','Complete joining/interface analysis, propose improvement, peer review, revise, and submit.')]:
        c.setFont('Helvetica-Bold',8); c.drawString(M+.1*inch,y,day+':'); wtext(c,desc,M+.75*inch,y,W-2*M-.85*inch,7.7,9.2); y-=.38*inch
    y=bar(c,y,'MATERIALS','One assembly per team is ideal')
    yy=y
    for i,m in enumerate(['student packet','purpose-made assemblies','rulers and calipers','plain/isometric graph paper','pencils and straightedges','optional: spare fasteners or alternate joints']):
        x=M+.15*inch+(i%2)*3.4*inch
        if i%2==0 and i>0: yy-=.24*inch
        c.setFont('Helvetica',7.8); c.drawString(x,yy,'- '+m)
    c.showPage()
    header(c,'IED 1.8 | Teacher Guide','Assessment and sample feedback',2)
    y=bar(c,H-1.02*inch,'ASSESSMENT LOOK-FORS','Use this to grade efficiently')
    for k,v in [('Sketching/views','Views are understandable, aligned, and use object, hidden, and center lines correctly.'),('Dimensions','Key sizes and feature locations are present. Students avoid repeated dimensions.'),('Measurement','Reported values are realistic for the tool used and include units.'),('Tolerances','Chosen tolerances connect to fit, alignment, function, or manufacturing.'),('Interface analysis','Students identify both parts in a joint and explain what the connection must accomplish.'),('Improvement','The proposed change improves a documented issue or tradeoff.')]:
        c.setFont('Helvetica-Bold',7.8); c.drawString(M+.1*inch,y,k+':'); wtext(c,v,M+1.45*inch,y,W-2*M-1.55*inch,7.6,9.2); y-=.35*inch
    y=bar(c,y-.05*inch,'COMMON STUDENT ISSUES','Use for quick feedback')
    for issue in ['Missing location dimensions for holes or slots','Only drawing one view when another view is needed','Repeating the same dimension in multiple places','Writing tolerance values without explaining why they matter','Calling every connection strong without addressing serviceability or assembly','Suggesting an improvement that is not connected to evidence']:
        c.setFont('Helvetica',7.6); c.drawString(M+.12*inch,y,'- '+issue); y-=.24*inch
    y=bar(c,y-.05*inch,'SAMPLE FEEDBACK PHRASES')
    for f in ['Your views communicate the shape well, but the hole/slot still needs a location dimension.','Your measurement record is strong; now connect one tolerance to a specific fit or alignment requirement.','Your joining analysis identifies the hardware, but it also needs to explain what the interface must do.','Your improvement is realistic. Add evidence from your measurements or interface analysis to support it.']:
        y=wtext(c,'- '+f,M+.12*inch,y,W-2*M-.22*inch,7.5,9.4)-.08*inch
    c.save()

def make_pptx():
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    colors={'navy':RGBColor(16,25,37),'blue':RGBColor(43,92,138),'green':RGBColor(47,125,105),'gold':RGBColor(243,201,92),'light':RGBColor(242,244,247),'white':RGBColor(255,255,255),'ink':RGBColor(17,17,17)}
    def slide(title,subtitle='',dark=False):
        s=prs.slides.add_slide(prs.slide_layouts[6])
        if dark:
            s.background.fill.solid(); s.background.fill.fore_color.rgb=colors['navy']
        t=s.shapes.add_textbox(Inches(.6),Inches(.42),Inches(10.2),Inches(.6)); p=t.text_frame.paragraphs[0]; r=p.add_run(); r.text=title; r.font.size=Pt(30); r.font.bold=True; r.font.color.rgb=colors['white'] if dark else colors['navy']
        if subtitle:
            st=s.shapes.add_textbox(Inches(.62),Inches(1.03),Inches(10),Inches(.3)); p=st.text_frame.paragraphs[0]; r=p.add_run(); r.text=subtitle; r.font.size=Pt(14); r.font.color.rgb=RGBColor(230,235,245) if dark else RGBColor(68,68,68)
        return s
    def card(s,x,y,w,h,title,body,accent='blue'):
        sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=colors['white']; sh.line.color.rgb=RGBColor(204,210,218)
        bar= s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x+.15),Inches(y+.22),Inches(.09),Inches(h-.44)); bar.fill.solid(); bar.fill.fore_color.rgb=colors[accent]; bar.line.color.rgb=colors[accent]
        tx=s.shapes.add_textbox(Inches(x+.35),Inches(y+.20),Inches(w-.5),Inches(h-.3)); tf=tx.text_frame; tf.word_wrap=True
        p=tf.paragraphs[0]; r=p.add_run(); r.text=title; r.font.bold=True; r.font.size=Pt(13); r.font.color.rgb=colors['navy']
        p=tf.add_paragraph(); r=p.add_run(); r.text=body; r.font.size=Pt(10); r.font.color.rgb=colors['ink']
    # slides
    s=slide('Engineering Documentation Challenge','Can another engineering team reproduce your assembly without guessing?',True)
    s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(1.0),Inches(4.9),Inches(2.4),Inches(.35)).fill.fore_color.rgb=colors['gold']
    for x,y,w,h,c in [(1.2,3.6,.55,1.2,'white'),(2.4,4.05,.5,.8,'light')]:
        sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=colors[c]; sh.line.color.rgb=colors[c]
    for word,x in [('SKETCH',4.3),('MEASURE',5.7),('DOCUMENT',7.2),('IMPROVE',9.1)]: card(s,x,4.25,1.15,.45,word,'', 'green')
    titles=[('Your Mission','Document an assembly from physical evidence.'),('What Strong Documentation Includes','Observe, sketch, measure, dimension, analyze, and improve.'),('Start With the Assembly','Identify parts before drawing details.'),('Sketching and View Planning','Choose views that communicate geometry.'),('Measurement and Dimensioning','Your dimensions should remove guessing.'),('Interface and Joining Analysis','Document how the assembly works together.'),('Engineering Improvement Proposal','Improve something real, not just something decorative.'),('Peer Review Standard','Could another team reproduce this without guessing?'),('Final Submission','Turn in one complete engineering documentation package.')]
    for i,(t,sub) in enumerate(titles,2):
        s=slide(t,sub)
        if i==2:
            card(s,.8,1.7,3.3,3.5,'The assembly is evidence','Observe, measure, sketch, and decide what another team needs to know.','blue'); card(s,4.8,1.7,3.3,3.5,'The package stands alone','A complete documentation package should make sense without verbal explanation.','green'); card(s,8.8,1.7,3.3,3.5,'The improvement is justified','Use evidence to support your design change.','blue')
        elif i==3:
            for j,(a,b) in enumerate([('Observe','Parts and interfaces'),('Sketch','Shape and orientation'),('Measure','Actual values'),('Dimension','Reproduce without guessing'),('Improve','Evidence-based change')]): card(s,.8+j*2.45,2.0,1.8,2.1,a,b, 'green' if j%2 else 'blue')
        elif i==4:
            card(s,5.2,1.4,5.8,1.1,'Part identification','Give each major component a simple part ID.','blue'); card(s,5.2,3.0,5.8,1.1,'Interface identification','Find where parts connect, align, or transfer load.','green'); card(s,5.2,4.6,5.8,1.1,'Functional purpose','Explain what the assembly is meant to accomplish.','gold')
        elif i==5:
            card(s,6.2,1.5,5.2,1.1,'Use aligned views','Features should line up between related views whenever possible.','blue'); card(s,6.2,3.0,5.2,1.1,'Use line conventions','Object, hidden, and center lines communicate visible and hidden features.','green'); card(s,6.2,4.5,5.2,1.1,'Choose useful views','Choose views that reveal needed geometry.','gold')
        elif i==6:
            card(s,5.2,1.4,5.8,1.1,'Measure first','Record real measurements before making final documentation decisions.','blue'); card(s,5.2,3.0,5.8,1.1,'Size and locate features','A feature must be sized and located.','green'); card(s,5.2,4.6,5.8,1.1,'Control critical dimensions','Use tolerances where function or fit depends on them.','gold')
        elif i==7:
            card(s,.8,1.6,3.4,1.4,'What joins the parts?','Identify the joining method and parts involved.','blue'); card(s,4.9,1.6,3.4,1.4,'What does it need to do?','Hold, align, transfer load, or allow service.','green'); card(s,9.0,1.6,3.4,1.4,'What is the tradeoff?','Every method has benefits and limitations.','gold')
            for j,word in enumerate(['LOAD','SERVICE','ALIGNMENT','MANUFACTURING','PERMANENCE']): card(s,.9+j*2.35,4.5,1.6,.5,word,'','blue')
        elif i==8:
            for j,(a,b) in enumerate([('Observation','What problem did you notice?'),('Evidence','What measurement or analysis supports it?'),('Change','What exactly would you modify?'),('Reason','How does the change improve the design?')]): card(s,1.0+j*3.0,2.1,2.4,2.1,a,b,'green' if j%2 else 'blue')
        elif i==9:
            card(s,1.0,2.6,3.2,2.0,'Find missing information','Look for unclear dimensions, line types, labels, or explanations.','blue'); card(s,5.1,2.6,3.2,2.0,'Give specific feedback','Write one strength and one revision that improves usability.','green'); card(s,9.2,2.6,3.2,2.0,'Revise before submitting','Peer review only matters if the package improves afterward.','gold')
        elif i==10:
            for j,(pts,cat) in enumerate([('20','Sketching and views'),('20','Dimensioning'),('20','Measurement and tolerances'),('15','Joining/interface analysis'),('15','Improvement proposal'),('10','Professional quality')]): card(s,1+(j%3)*3.8,1.6+(j//3)*1.8,3.1,1.0,pts+' pts',cat,'green' if j%2 else 'blue')
    prs.save(str(PPTX))

make_student(); make_teacher(); make_pptx()
print('Built IED 1.8 resources')
